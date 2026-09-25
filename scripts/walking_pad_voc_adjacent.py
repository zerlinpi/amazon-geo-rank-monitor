#!/usr/bin/env python3
import csv,json,os,re,hashlib,gzip,io
from collections import Counter
from datetime import datetime,timezone
from pathlib import Path
import requests

TARGET=int(os.getenv("TARGET_REVIEWS","40000"))
OUTDIR=Path(os.getenv("OUTDIR","adjacent_output")); OUTDIR.mkdir(parents=True,exist_ok=True)
META_URL="https://mcauleylab.ucsd.edu/public_datasets/data/amazon_2023/raw/meta_categories/meta_Sports_and_Outdoors.jsonl.gz"
REV_URL="https://mcauleylab.ucsd.edu/public_datasets/data/amazon_2023/raw/review_categories/Sports_and_Outdoors.jsonl.gz"

STRICT=[
"walking pad","walkingpad","under desk treadmill","under-desk treadmill","desk treadmill","walking treadmill",
"portable treadmill","compact treadmill","foldable treadmill","folding treadmill","mini treadmill",
"2 in 1 treadmill","2-in-1 treadmill","home office treadmill","treadmill for office","treadmill under desk","walking machine"
]
DESCRIPTORS=["walking","under desk","desk","office","workstation","folding","foldable","compact","portable","mini",
"small","slim","space saving","space-saving","low profile","2 in 1","2-in-1","home"]
CATS={
"安全/召回":r"recall|fire|burn|smoke|unsafe|fall|fell|injur|shock|sudden stop|abrupt stop|almost fell",
"耐久/质量":r"fail|broke|broken|stopped working|died|motor|overheat|hot|burnt|squeak|grind|lasted",
"跑带/稳定性":r"belt.*slip|slipping|belt.*drift|belt.*shift|center.*belt|align|wobbl|unstable|jerk",
"噪音/振动":r"noisy|loud|noise|quiet|vibrat|stomp|downstairs|neighbor|creak",
"尺寸/适配":r"too short|too narrow|wider|longer|stride|tall|height|width|capacity|weight limit",
"收纳/便携":r"fold|stor(e|age)|under the bed|under bed|under sofa|under couch|upright|vertical|heavy|wheel",
"App/遥控/数据":r"app|bluetooth|remote|disconnect|apple health|apple watch|subscription",
"售后/保修":r"warranty|refund|return|customer service|support|replacement|parts|repair",
"维护/跑带":r"lubricat|maintenance|clean|tension|calibrat|adjust",
"速度/坡度/参数":r"incline|speed|mph|km/h|slows|slowing|faster",
"人体工学/办公":r"handle|work from home|wfh|desk|meeting|zoom|typing|phone holder",
"价格/价值":r"expensive|overpriced|price|worth|cheap|value",
"运输/到货":r"arrived damaged|damaged|shipping|delivery|missing|box",
}
FIELDS=["platform","amazon_category","relevance_tier","source_type","source_url","date","rating","verified_purchase",
"helpful_vote","brand","model","asin","parent_asin","review_title","comment","category","dedup_key"]

def norm(s):return re.sub(r"\s+"," ",str(s or "").strip().lower())
def text(v):
    if v is None:return ""
    if isinstance(v,(list,tuple)):return " ".join(map(str,v))
    if isinstance(v,dict):return json.dumps(v,ensure_ascii=False)
    return str(v)
def strict(blob):return any(k in blob for k in STRICT)
def broad(blob):
    return "treadmill" in blob and any(d in blob for d in DESCRIPTORS)
def classify(s):
    low=norm(s);best=("其他/待聚类",0)
    for k,p in CATS.items():
        n=len(re.findall(p,low,re.I))
        if n>best[1]:best=(k,n)
    return best[0]
def dk(t,p):return hashlib.sha256((norm(t)+"|"+norm(p)).encode()).hexdigest()
def dt(ts):
    try:
        x=float(ts);x=x/1000 if x>10**12 else x
        return datetime.fromtimestamp(x,tz=timezone.utc).date().isoformat()
    except:return ""
def it(url,label):
    print("Streaming",label,flush=True)
    with requests.get(url,stream=True,timeout=(30,1200),headers={"User-Agent":"walking-pad-voc-research/5.0"}) as r:
        r.raise_for_status();r.raw.decode_content=False
        with gzip.GzipFile(fileobj=r.raw) as gz:
            f=io.TextIOWrapper(gz,encoding="utf-8",errors="replace")
            for line in f:
                try:yield json.loads(line)
                except:continue

products={};scanned=0
for row in it(META_URL,"2023 adjacent metadata"):
    scanned+=1
    blob=norm(" ".join([text(row.get("title")),text(row.get("description")),text(row.get("features")),
                        text(row.get("categories")),text(row.get("store"))]))
    if broad(blob) and not strict(blob):
        p=str(row.get("parent_asin") or "")
        if p:products[p]={"title":text(row.get("title")),"brand":text(row.get("store"))}
    if scanned%100000==0:print(f"meta={scanned:,}; adjacent products={len(products):,}",flush=True)
print(f"metadata done {scanned:,}; adjacent products={len(products):,}",flush=True)

rows=[];seen=set();cats=Counter();rscan=raw=0
for rv in it(REV_URL,"2023 adjacent reviews"):
    rscan+=1;p=str(rv.get("parent_asin") or "")
    if p not in products:
        if rscan%500000==0:print(f"reviews={rscan:,}; kept={len(rows):,}",flush=True)
        continue
    body=text(rv.get("text")).strip()
    if len(body)<8:continue
    raw+=1;key=dk(body,p)
    if key in seen:continue
    seen.add(key);prod=products[p];title=text(rv.get("title"))
    rows.append({"platform":"Amazon Reviews 2023","amazon_category":"Sports_and_Outdoors",
    "relevance_tier":"Adjacent Compact/Folding Treadmill","source_type":"historical adjacent product review",
    "source_url":f"https://www.amazon.com/dp/{p}","date":dt(rv.get("timestamp")),"rating":rv.get("rating"),
    "verified_purchase":rv.get("verified_purchase"),"helpful_vote":rv.get("helpful_vote"),"brand":prod["brand"],
    "model":prod["title"],"asin":text(rv.get("asin")),"parent_asin":p,"review_title":title,"comment":body,
    "category":classify(title+" "+body),"dedup_key":key})
    cats[rows[-1]["category"]]+=1
    if len(rows)>=TARGET:break
    if rscan%500000==0:print(f"reviews={rscan:,}; raw={raw:,}; kept={len(rows):,}",flush=True)

with open(OUTDIR/"walking_pad_adjacent_2023.csv","w",encoding="utf-8-sig",newline="") as f:
    w=csv.DictWriter(f,fieldnames=FIELDS);w.writeheader();w.writerows(rows)
from openpyxl import Workbook
wb=Workbook(write_only=True);ws=wb.create_sheet("评论明细_Adjacent");ws.append(FIELDS)
for r in rows:ws.append([r.get(h,"") for h in FIELDS])
sw=wb.create_sheet("汇总");sw.append(["指标","值"]);sw.append(["去重评论",len(rows)]);sw.append(["目标",TARGET]);sw.append(["匹配商品",len(products)])
for k,v in cats.most_common():sw.append([k,v])
wb.save(OUTDIR/"walking_pad_adjacent_2023.xlsx")
stats={"target":TARGET,"final_rows":len(rows),"matched_products":len(products),"reviews_scanned":rscan,"raw_matches":raw,"pain_categories":dict(cats)}
with open(OUTDIR/"stats.json","w",encoding="utf-8") as f:json.dump(stats,f,ensure_ascii=False,indent=2)
print(json.dumps(stats,ensure_ascii=False),flush=True)
