#!/usr/bin/env python3
import csv, json, os, re, hashlib, gzip, io
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
import requests

TARGET=int(os.getenv("TARGET_REVIEWS","50000"))
OUTDIR=Path(os.getenv("OUTDIR","voc2018_output")); OUTDIR.mkdir(parents=True,exist_ok=True)
BASE="https://mcauleylab.ucsd.edu/public_datasets/data/amazon_v2"
CATEGORIES=["Sports_and_Outdoors","Home_and_Kitchen","Office_Products","Tools_and_Home_Improvement"]

KEYWORDS=[
    "walking pad","walkingpad","under desk treadmill","under-desk treadmill",
    "desk treadmill","walking treadmill","portable treadmill","compact treadmill",
    "foldable treadmill","folding treadmill","mini treadmill","2 in 1 treadmill",
    "2-in-1 treadmill","home office treadmill","treadmill for office",
    "treadmill under desk","walking machine","folding walking treadmill",
    "portable walking treadmill","workstation treadmill","treadmill desk"
]
BRANDS=[
    "walkingpad","kingsmith","urevo","deerrun","sperax","egofit","goyouth",
    "goplus","merach","lifespan","sunny health","toputure","maksone","wellfit",
    "axefit","motiongrey","freepi","vitalwalk","rebel treadmill","imovr"
]
CATEGORY_PATTERNS={
    "安全/召回":r"recall|fire|burn|smoke|unsafe|fall|fell|injur|shock|sudden stop|abrupt stop|almost fell",
    "耐久/质量":r"fail|broke|broken|stopped working|died|motor|overheat|hot|burnt|squeak|grind|lasted",
    "跑带/稳定性":r"belt.*slip|slipping|belt.*drift|belt.*shift|center.*belt|align|wobbl|unstable|jerk",
    "噪音/振动":r"noisy|loud|noise|quiet|vibrat|stomp|downstairs|neighbor|creak",
    "尺寸/适配":r"too short|too narrow|wider|longer|stride|tall|height|width|capacity|weight limit",
    "收纳/便携":r"fold|stor(e|age)|under the bed|under bed|under sofa|under couch|upright|vertical|heavy|wheel",
    "App/遥控/数据":r"app|bluetooth|remote|disconnect|apple health|apple watch|subscription",
    "售后/保修":r"warranty|refund|return|customer service|support|replacement|parts|repair",
    "维护/跑带":r"lubricat|maintenance|clean|tension|calibrat|adjust",
    "速度/坡度/参数":r"incline|speed|mph|km/h|slows|slowing|faster",
    "人体工学/办公":r"handle|work from home|wfh|desk|meeting|zoom|typing|phone holder",
    "价格/价值":r"expensive|overpriced|price|worth|cheap|value",
    "运输/到货":r"arrived damaged|damaged|shipping|delivery|missing|box",
}
FIELDS=["platform","amazon_category","source_type","source_url","date","rating","verified_purchase","helpful_vote",
        "brand","model","asin","parent_asin","review_title","comment","category","dedup_key"]

def norm(s): return re.sub(r"\s+"," ",str(s or "").strip().lower())
def as_text(v):
    if v is None:return ""
    if isinstance(v,(list,tuple)):return " ".join(map(str,v))
    if isinstance(v,dict):return json.dumps(v,ensure_ascii=False)
    return str(v)
def match_product(row):
    blob=norm(" ".join([
        as_text(row.get("title")),as_text(row.get("description")),as_text(row.get("feature")),
        as_text(row.get("categories")),as_text(row.get("brand")),as_text(row.get("tech1")),as_text(row.get("tech2"))
    ]))
    if any(k in blob for k in KEYWORDS): return True
    return any(b in blob for b in BRANDS) and "treadmill" in blob and any(
        x in blob for x in ["walking","under desk","foldable","folding","compact","portable","desk","workstation","mini"])
def classify(text):
    low=norm(text); best=("其他/待聚类",0)
    for cat,pat in CATEGORY_PATTERNS.items():
        n=len(re.findall(pat,low,flags=re.I))
        if n>best[1]: best=(cat,n)
    return best[0]
def dkey(text, asin): return hashlib.sha256((norm(text)+"|"+norm(asin)).encode()).hexdigest()
def iso_date(ts):
    try:return datetime.fromtimestamp(float(ts),tz=timezone.utc).date().isoformat()
    except:return ""
def iter_gz(url,label):
    print(f"Streaming {label}: {url}",flush=True)
    with requests.get(url,stream=True,timeout=(30,1200),headers={"User-Agent":"walking-pad-voc-research/4.0"}) as r:
        r.raise_for_status(); r.raw.decode_content=False
        with gzip.GzipFile(fileobj=r.raw,mode="rb") as gz:
            txt=io.TextIOWrapper(gz,encoding="utf-8",errors="replace")
            for line in txt:
                try: yield json.loads(line)
                except: continue

rows=[]; seen=set(); cats=Counter(); product_rows=[]; stats_by_cat={}
for amazon_cat in CATEGORIES:
    if len(rows)>=TARGET: break
    meta_url=f"{BASE}/metaFiles2/meta_{amazon_cat}.json.gz"
    rev_url=f"{BASE}/categoryFiles/{amazon_cat}.json.gz"
    products={}; meta_scanned=0
    try:
        for row in iter_gz(meta_url,f"2018 {amazon_cat} metadata"):
            meta_scanned+=1
            if match_product(row):
                asin=str(row.get("asin") or "")
                if asin:
                    products[asin]={
                        "asin":asin,"title":as_text(row.get("title")),"brand":as_text(row.get("brand")),
                        "price":row.get("price"),"amazon_category":amazon_cat
                    }
            if meta_scanned%250000==0:
                print(f"{amazon_cat} metadata={meta_scanned:,}; matched={len(products):,}",flush=True)
    except requests.HTTPError as e:
        print(f"SKIP metadata {amazon_cat}: {e}",flush=True); continue
    print(f"{amazon_cat} metadata complete: {meta_scanned:,}; matched={len(products):,}",flush=True)
    product_rows.extend(products.values())
    scanned=raw=added=0
    try:
        for rv in iter_gz(rev_url,f"2018 {amazon_cat} reviews"):
            scanned+=1
            asin=str(rv.get("asin") or "")
            if asin not in products:
                if scanned%1000000==0:
                    print(f"{amazon_cat} reviews={scanned:,}; total={len(rows):,}",flush=True)
                continue
            text=as_text(rv.get("reviewText")).strip()
            if len(text)<8: continue
            raw+=1
            dk=dkey(text,asin)
            if dk in seen: continue
            seen.add(dk)
            p=products[asin]; title=as_text(rv.get("summary"))
            row={
                "platform":"Amazon Reviews 2018","amazon_category":amazon_cat,
                "source_type":"historical product review","source_url":f"https://www.amazon.com/dp/{asin}",
                "date":iso_date(rv.get("unixReviewTime")),"rating":rv.get("overall"),
                "verified_purchase":rv.get("verified"),"helpful_vote":rv.get("vote"),
                "brand":p["brand"],"model":p["title"],"asin":asin,"parent_asin":asin,
                "review_title":title,"comment":text,"category":classify(title+" "+text),"dedup_key":dk
            }
            rows.append(row); cats[row["category"]]+=1; added+=1
            if len(rows)>=TARGET: break
            if scanned%1000000==0:
                print(f"{amazon_cat} reviews={scanned:,}; raw={raw:,}; added={added:,}; total={len(rows):,}",flush=True)
    except requests.HTTPError as e:
        print(f"SKIP reviews {amazon_cat}: {e}",flush=True)
    stats_by_cat[amazon_cat]={"meta_scanned":meta_scanned,"matched_products":len(products),"reviews_scanned":scanned,"raw_matches":raw,"dedup_added":added}
    print(f"{amazon_cat} done: scanned={scanned:,}; raw={raw:,}; added={added:,}; total={len(rows):,}",flush=True)

with open(OUTDIR/"walking_pad_voc_2018.csv","w",encoding="utf-8-sig",newline="") as f:
    w=csv.DictWriter(f,fieldnames=FIELDS); w.writeheader(); w.writerows(rows)
with open(OUTDIR/"matched_products_2018.csv","w",encoding="utf-8-sig",newline="") as f:
    pf=["amazon_category","asin","title","brand","price"]; w=csv.DictWriter(f,fieldnames=pf); w.writeheader()
    for p in product_rows:w.writerow({k:p.get(k) for k in pf})
with open(OUTDIR/"summary_2018.csv","w",encoding="utf-8-sig",newline="") as f:
    w=csv.writer(f);w.writerow(["category","count"]);w.writerows(cats.most_common())

from openpyxl import Workbook
wb=Workbook(write_only=True)
ws=wb.create_sheet("评论明细_2018"); ws.append(FIELDS)
for r in rows: ws.append([r.get(h,"") for h in FIELDS])
sw=wb.create_sheet("汇总");sw.append(["指标","值"]);sw.append(["去重评论",len(rows)]);sw.append(["目标",TARGET])
for cat,n in cats.most_common(): sw.append([cat,n])
wb.save(OUTDIR/"walking_pad_voc_2018.xlsx")

stats={"target":TARGET,"final_rows":len(rows),"categories":stats_by_cat,"pain_categories":dict(cats),"matched_products":len(product_rows)}
with open(OUTDIR/"stats.json","w",encoding="utf-8") as f:json.dump(stats,f,ensure_ascii=False,indent=2)
print(json.dumps(stats,ensure_ascii=False),flush=True)
