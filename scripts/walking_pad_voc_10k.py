#!/usr/bin/env python3
import csv, json, os, re, sys, time, hashlib
from collections import Counter
from datetime import datetime, timezone

TARGET = int(os.getenv("TARGET_REVIEWS", "12000"))
MIN_VALID = int(os.getenv("MIN_VALID", "10000"))
OUTDIR = os.getenv("OUTDIR", "voc_output")
os.makedirs(OUTDIR, exist_ok=True)

KEYWORDS = [
    "walking pad","walkingpad","under desk treadmill","under-desk treadmill",
    "desk treadmill","walking treadmill","portable treadmill","compact treadmill",
    "foldable treadmill","folding treadmill","mini treadmill","2 in 1 treadmill",
    "2-in-1 treadmill","home office treadmill","treadmill for office"
]
BRANDS = [
    "walkingpad","kingsmith","urevo","deerrun","sperax","egofit","goyouth",
    "goplus","merach","lifespan","sunny health","toputure","maksone","wellfit",
    "axefit","motiongrey","freepi","vitalwalk"
]
NEG_PATTERNS = {
    "安全/召回": r"recall|fire|burn|smoke|unsafe|fall|fell|injur|shock|sudden stop|abrupt stop",
    "耐久/质量": r"fail|broke|broken|stopped working|died|motor|overheat|hot|burnt|squeak|grind",
    "跑带/稳定性": r"belt.*slip|slipping|belt.*drift|belt.*shift|align|wobbl|unstable|jerk",
    "噪音/振动": r"noisy|loud|noise|quiet|vibrat|stomp|downstairs|neighbor|creak",
    "尺寸/适配": r"too short|too narrow|wider|longer|stride|tall|height|width|capacity|weight limit",
    "收纳/便携": r"fold|stor(e|age)|under the bed|under bed|under sofa|under couch|upright|vertical|heavy|wheel",
    "App/遥控/数据": r"app|bluetooth|remote|disconnect|apple health|apple watch|subscription",
    "售后/保修": r"warranty|refund|return|customer service|support|replacement|parts|repair",
    "维护/跑带": r"lubricat|maintenance|clean|tension|calibrat|adjust",
    "速度/坡度/参数": r"incline|speed|mph|km/h|slows|slowing|faster",
    "人体工学/办公": r"handle|work from home|wfh|desk|meeting|zoom|typing|phone holder",
    "价格/价值": r"expensive|overpriced|price|worth|cheap|value",
    "运输/到货": r"arrived damaged|damaged|shipping|delivery|missing|box",
}

def norm(s):
    return re.sub(r"\s+", " ", (s or "").strip().lower())

def match_product(item):
    title = str(item.get("title") or "")
    desc = item.get("description") or []
    feats = item.get("features") or []
    cats = item.get("categories") or []
    store = str(item.get("store") or "")
    if isinstance(desc, list): desc = " ".join(map(str, desc))
    if isinstance(feats, list): feats = " ".join(map(str, feats))
    if isinstance(cats, list): cats = " ".join(map(str, cats))
    blob = norm(" ".join([title, str(desc), str(feats), str(cats), store]))
    if any(k in blob for k in KEYWORDS):
        return True
    if "treadmill" in blob and any(b in blob for b in BRANDS):
        return True
    return False

def classify(text):
    low = norm(text)
    best = ("其他/待聚类", 0)
    for cat, pat in NEG_PATTERNS.items():
        n = len(re.findall(pat, low, re.I))
        if n > best[1]:
            best = (cat, n)
    return best[0]

def dedup_key(text, parent):
    return hashlib.sha256((norm(text)+"|"+parent).encode("utf-8")).hexdigest()

def iso_date(ts):
    try:
        ts = int(ts)
        if ts > 10**12:
            ts = ts/1000
        return datetime.fromtimestamp(ts, tz=timezone.utc).date().isoformat()
    except Exception:
        return ""

print("Installing/initializing datasets...", flush=True)
from datasets import load_dataset

meta_cfg = "raw_meta_Sports_and_Outdoors"
review_cfg = "raw_review_Sports_and_Outdoors"

products = {}
print("Scanning Sports & Outdoors metadata for walking-pad/treadmill products...", flush=True)
meta = load_dataset("McAuley-Lab/Amazon-Reviews-2023", meta_cfg, split="full", streaming=True, trust_remote_code=True)
for i, item in enumerate(meta, 1):
    if match_product(item):
        p = str(item.get("parent_asin") or "")
        if p:
            products[p] = {
                "parent_asin": p,
                "title": str(item.get("title") or ""),
                "brand": str(item.get("store") or ""),
                "rating_number": item.get("rating_number"),
                "average_rating": item.get("average_rating"),
                "price": item.get("price"),
            }
    if i % 100000 == 0:
        print(f"metadata scanned={i:,}, matched products={len(products):,}", flush=True)

print(f"Matched products: {len(products):,}", flush=True)
with open(os.path.join(OUTDIR, "matched_products.csv"), "w", encoding="utf-8-sig", newline="") as f:
    w = csv.DictWriter(f, fieldnames=["parent_asin","title","brand","rating_number","average_rating","price"])
    w.writeheader(); w.writerows(products.values())

if not products:
    raise SystemExit("No matching products found.")

fields = [
    "platform","source_type","source_url","date","rating","verified_purchase","helpful_vote",
    "brand","model","asin","parent_asin","review_title","comment","category","dedup_key"
]
raw_path = os.path.join(OUTDIR, "walking_pad_reviews_raw.csv")
dedup_path = os.path.join(OUTDIR, "walking_pad_reviews_10k.csv")
seen = set()
valid = []
raw_count = 0
scanned = 0
cats = Counter()

print("Streaming reviews...", flush=True)
reviews = load_dataset("McAuley-Lab/Amazon-Reviews-2023", review_cfg, split="full", streaming=True, trust_remote_code=True)
with open(raw_path, "w", encoding="utf-8-sig", newline="") as f:
    w = csv.DictWriter(f, fieldnames=fields); w.writeheader()
    for rv in reviews:
        scanned += 1
        parent = str(rv.get("parent_asin") or "")
        if parent not in products:
            if scanned % 500000 == 0:
                print(f"reviews scanned={scanned:,}, raw matches={raw_count:,}, valid dedup={len(valid):,}", flush=True)
            continue
        text = str(rv.get("text") or "").strip()
        if len(text) < 8:
            continue
        title = str(rv.get("title") or "")
        product = products[parent]
        cat = classify(title + " " + text)
        dk = dedup_key(text, parent)
        row = {
            "platform":"Amazon Reviews 2023",
            "source_type":"historical review dataset",
            "source_url":f"https://www.amazon.com/dp/{parent}",
            "date":iso_date(rv.get("timestamp")),
            "rating":rv.get("rating"),
            "verified_purchase":rv.get("verified_purchase"),
            "helpful_vote":rv.get("helpful_vote"),
            "brand":product["brand"],
            "model":product["title"],
            "asin":str(rv.get("asin") or ""),
            "parent_asin":parent,
            "review_title":title,
            "comment":text,
            "category":cat,
            "dedup_key":dk
        }
        w.writerow(row)
        raw_count += 1
        if dk not in seen:
            seen.add(dk)
            valid.append(row)
            cats[cat] += 1
        if len(valid) >= TARGET:
            break
        if scanned % 500000 == 0:
            print(f"reviews scanned={scanned:,}, raw matches={raw_count:,}, valid dedup={len(valid):,}", flush=True)

print(f"Finished review stream: scanned={scanned:,}, raw={raw_count:,}, dedup={len(valid):,}", flush=True)
if len(valid) < MIN_VALID:
    print(f"WARNING: only {len(valid):,} deduplicated reviews found (< {MIN_VALID:,}).", flush=True)

with open(dedup_path, "w", encoding="utf-8-sig", newline="") as f:
    w = csv.DictWriter(f, fieldnames=fields); w.writeheader(); w.writerows(valid)

# XLSX
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

xlsx = os.path.join(OUTDIR, "walking_pad_voc_10k.xlsx")
wb = Workbook()
ws = wb.active
ws.title = "评论明细_All"
headers = fields
ws.append(headers)
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="121826")
    c.alignment = Alignment(vertical="center", wrap_text=True)
for r in valid:
    ws.append([r[h] for h in headers])
ws.freeze_panes = "A2"
widths = [18,22,36,13,9,16,12,20,45,16,16,30,80,18,68]
for i,wid in enumerate(widths,1):
    ws.column_dimensions[get_column_letter(i)].width = wid
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(vertical="top", wrap_text=True)

sumws = wb.create_sheet("汇总")
sumws.append(["指标","值"])
sumws.append(["实际去重评论",len(valid)])
sumws.append(["原始匹配评论",raw_count])
sumws.append(["扫描评论总数",scanned])
sumws.append(["匹配商品数",len(products)])
sumws.append(["目标",TARGET])
sumws.append(["生成时间",datetime.now(timezone.utc).isoformat()])
sumws.append([])
sumws.append(["痛点分类","样本数"])
for cat,n in cats.most_common():
    sumws.append([cat,n])
for c in sumws[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="121826")
sumws.column_dimensions["A"].width = 30
sumws.column_dimensions["B"].width = 20

prodws = wb.create_sheet("匹配商品")
prodws.append(["parent_asin","title","brand","rating_number","average_rating","price"])
for p in products.values():
    prodws.append([p[k] for k in ["parent_asin","title","brand","rating_number","average_rating","price"]])
for c in prodws[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="121826")
prodws.freeze_panes = "A2"
prodws.column_dimensions["A"].width=18
prodws.column_dimensions["B"].width=60
prodws.column_dimensions["C"].width=24

wb.save(xlsx)

stats = {
    "target":TARGET,
    "min_valid":MIN_VALID,
    "matched_products":len(products),
    "reviews_scanned":scanned,
    "raw_matches":raw_count,
    "dedup_valid":len(valid),
    "categories":dict(cats)
}
with open(os.path.join(OUTDIR,"stats.json"),"w",encoding="utf-8") as f:
    json.dump(stats,f,ensure_ascii=False,indent=2)
print(json.dumps(stats, ensure_ascii=False), flush=True)
